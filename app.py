# -*- coding: utf-8 -*-
"""
KBLT Converter - Web app chuyển đổi file Excel gốc sang file mẫu KBLT
Hỗ trợ deploy cloud (Render) và chạy local
"""

import os
import uuid
import unicodedata
import threading
import time
from datetime import datetime, timedelta
from flask import Flask, render_template, request, jsonify, send_file

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = os.path.join(os.path.dirname(__file__), 'uploads')
app.config['OUTPUT_FOLDER'] = os.path.join(os.path.dirname(__file__), 'outputs')
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024  # 50MB max

os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
os.makedirs(app.config['OUTPUT_FOLDER'], exist_ok=True)

# File mẫu nằm trong project (relative path)
TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), 'data', 'tblt_vn_import.xlsx')

# Tự động dọn file output cũ hơn 1 giờ
def cleanup_old_files():
    while True:
        time.sleep(600)  # Mỗi 10 phút
        now = time.time()
        for folder in [app.config['UPLOAD_FOLDER'], app.config['OUTPUT_FOLDER']]:
            try:
                for f in os.listdir(folder):
                    fpath = os.path.join(folder, f)
                    if os.path.isfile(fpath) and now - os.path.getmtime(fpath) > 3600:
                        os.remove(fpath)
            except Exception:
                pass

cleanup_thread = threading.Thread(target=cleanup_old_files, daemon=True)
cleanup_thread.start()


def nfc(s):
    if s is None:
        return ''
    return unicodedata.normalize('NFC', str(s).strip())

def build_lookups(wb_tpl):
    lookups = {}
    ws_tt = wb_tpl['TINH_THANH']
    tinh_code = {}
    tinh_display = {}
    for row in ws_tt.iter_rows(min_row=2, values_only=True):
        if row[1] and row[2]:
            ma = str(row[0]).strip()
            ten = nfc(row[1])
            tinh_code[ten] = ma
            tinh_display[ten] = row[2]
    lookups['tinh_code'] = tinh_code
    lookups['tinh_display'] = tinh_display

    ws_px = wb_tpl['PHUONG_XA']
    px_map = {}
    for row in ws_px.iter_rows(min_row=2, values_only=True):
        if row[1] and row[2] and row[3]:
            ma_tinh = str(row[2]).strip()
            ten_px = nfc(row[1])
            if ma_tinh not in px_map:
                px_map[ma_tinh] = {}
            px_map[ma_tinh][ten_px] = row[3]
    lookups['px_map'] = px_map

    ws_dm = wb_tpl['DANH_MUC']
    loai_giay_to = {}
    noi_cu_tru = {}
    ly_do = {}
    gioi_tinh = {}
    quoc_tich = {}
    for row in ws_dm.iter_rows(min_row=2, values_only=True):
        if row[0]:
            val = nfc(row[0])
            loai_giay_to[val] = val
            parts = val.split(' - ')
            if len(parts) > 1:
                short = parts[1].strip()
                loai_giay_to[short] = val
                loai_giay_to[short.replace('Thẻ ', '')] = val
                if "Khác" in short or "khác" in short:
                    loai_giay_to['Giấy tờ khác'] = val
                elif "CCCD" in short:
                    loai_giay_to['CCCD'] = val
                elif "CMND" in short:
                    loai_giay_to['CMND'] = val

        if row[1]:
            val = nfc(row[1])
            noi_cu_tru[val] = val
            parts = val.split(' - ')
            if len(parts) > 1:
                noi_cu_tru[parts[1].strip()] = val
            if "Khác" in val:
                noi_cu_tru['Khác'] = val
                noi_cu_tru['Địa chỉ khác'] = val

        if row[2]:
            val = nfc(row[2])
            ly_do[val] = val
            parts = val.split(' - ')
            if len(parts) > 1:
                short = parts[1].strip()
                ly_do[short] = val
                if short == "Chữa bệnh":
                    ly_do["Khám chữa bệnh"] = val
                    ly_do["KCB"] = val

        if row[3]:
            val = nfc(row[3])
            gioi_tinh[val] = val
            parts = val.split(' - ')
            if len(parts) > 1:
                gioi_tinh[parts[0].strip()] = val
                gioi_tinh[parts[1].strip()] = val

        if row[4]:
            val = nfc(row[4])
            quoc_tich[val] = val
            parts = val.split(' - ')
            if len(parts) > 1:
                quoc_tich[parts[0].strip()] = val
            if "Viet Nam" in val or "Việt Nam" in val:
                quoc_tich["VNM"] = val

    lookups['loai_giay_to'] = loai_giay_to
    lookups['noi_cu_tru'] = noi_cu_tru
    lookups['ly_do'] = ly_do
    lookups['gioi_tinh'] = gioi_tinh
    lookups['quoc_tich'] = quoc_tich
    return lookups

def map_tinh(raw_value, lookups):
    if not raw_value:
        return ''
    val = nfc(raw_value)
    clean = val.replace('Tỉnh ', '').replace('Thành phố ', 'TP. ')
    if clean in lookups['tinh_display']:
        return lookups['tinh_display'][clean]
    clean2 = val.replace('Tỉnh ', '').replace('Thành phố ', '')
    for k, v in lookups['tinh_display'].items():
        if clean2 in k or k in clean2:
            return v
    return val

def get_tinh_code(raw_value, lookups):
    if not raw_value:
        return ''
    val = nfc(raw_value)
    clean = val.replace('Tỉnh ', '').replace('Thành phố ', 'TP. ')
    return lookups['tinh_code'].get(clean, '')

def map_phuong_xa(raw_px, tinh_code, lookups):
    if not raw_px or not tinh_code:
        return ''
    val = nfc(raw_px)
    if tinh_code in lookups['px_map'] and val in lookups['px_map'][tinh_code]:
        return lookups['px_map'][tinh_code][val]
    return val

def map_generic(raw_value, lookup_dict):
    if not raw_value:
        return ''
    val = nfc(raw_value)
    if val in lookup_dict:
        return lookup_dict[val]
    return val

def fix_date(raw_value):
    if not raw_value:
        return None, None
    val = str(raw_value).strip()
    date_part = val.split(' ')[0] if ' ' in val else val
    try:
        dt = datetime.strptime(date_part, '%d/%m/%Y')
        return date_part, dt
    except:
        return date_part, None

def find_source_columns(ws_src):
    """Dynamically find column indices by scanning rows for headers."""
    header_row_idx = None
    col_map = {}
    
    # Scan first 20 rows to find header row (looking for "Họ và tên")
    for row_idx in range(1, 21):
        for col_idx in range(1, ws_src.max_column + 1):
            val = nfc(ws_src.cell(row=row_idx, column=col_idx).value).lower()
            if 'họ và tên' in val:
                header_row_idx = row_idx
                break
        if header_row_idx:
            break
            
    if not header_row_idx:
        return None, None
        
    # Pre-read headers to disambiguate 'đến ngày' vs 'từ ngày' vs 'đi ngày'
    headers = []
    for col_idx in range(1, ws_src.max_column + 1):
        headers.append(nfc(ws_src.cell(row=header_row_idx, column=col_idx).value).lower())
    
    has_di_ngay = any('đi ngày' in h for h in headers if h)
    has_tu_ngay = any('từ ngày' in h for h in headers if h)

    # Map columns based on header texts
    for col_idx, val in enumerate(headers):
        if not val:
            continue
            
        if 'họ và tên' in val:
            col_map['ho_ten'] = col_idx
        elif 'ngày, tháng, năm sinh' in val or 'ngày sinh' in val:
            col_map['ngay_sinh'] = col_idx
        elif 'giới tính' in val:
            col_map['gioi_tinh'] = col_idx
        elif 'quốc tịch' in val:
            col_map['quoc_tich'] = col_idx
        elif 'loại giấy tờ' in val:
            col_map['loai_giay_to'] = col_idx
        elif 'tên giấy tờ' in val:
            col_map['ten_giay_to'] = col_idx
        elif 'số giấy tờ' in val:
            col_map['so_giay_to'] = col_idx
        elif 'cmnd' in val or 'cccd' in val or 'số định danh' in val:
            col_map['so_cccd'] = col_idx
        elif 'hộ chiếu' in val:
            col_map['so_ho_chieu'] = col_idx
        elif 'giấy tờ khác' in val:
            col_map['so_giay_to_khac'] = col_idx
        elif 'số điện thoại' in val:
            col_map['so_dt'] = col_idx
        elif 'loại cư trú' in val or 'nơi cư trú' in val:
            col_map['noi_cu_tru'] = col_idx
        elif 'tỉnh/tp' in val or 'tỉnh, thành' in val or 'tỉnh thành' in val:
            col_map['tinh_tp'] = col_idx
        elif 'phường/xã' in val or 'phường xã' in val:
            col_map['phuong_xa'] = col_idx
        elif 'địa chỉ chi tiết' in val or 'số nhà' in val:
            col_map['dia_chi'] = col_idx
        elif 'từ ngày' in val or 'ngày đến' in val:
            col_map['ngay_den'] = col_idx
        elif 'đi ngày' in val or 'ngày đi' in val:
            col_map['ngay_di'] = col_idx
        elif 'đến ngày' in val:
            if has_di_ngay:
                col_map['ngay_den'] = col_idx  # 'đến' = arrive
            else:
                col_map['ngay_di'] = col_idx   # 'đến' = depart
        elif 'lý do lưu trú' in val or 'lý do cư trú' in val:
            col_map['ly_do'] = col_idx
        elif 'phòng/khoa' in val or 'số phòng' in val:
            col_map['phong_khoa'] = col_idx

    return header_row_idx, col_map

def process_files(source_path, extra_days=3):
    import openpyxl
    
    stats = {
        'total_rows': 0, 'mapped_gender': 0, 'mapped_doc_type': 0,
        'mapped_residence': 0, 'mapped_province': 0, 'mapped_ward': 0,
        'mapped_nationality': 0, 'mapped_reason': 0, 'fixed_date_format': 0,
        'added_departure_date': 0, 'errors': [],
    }

    try:
        wb_tpl = openpyxl.load_workbook(TEMPLATE_PATH)
        wb_src = openpyxl.load_workbook(source_path)
    except Exception as e:
        stats['errors'].append(f'Lỗi đọc file: {str(e)}')
        return None, stats

    ws_tpl = wb_tpl['DS_KHACH_VIET_NAM_LUU_TRU']
    ws_src = wb_src.active

    lookups = build_lookups(wb_tpl)

    # Dynamic column mapping
    header_row_idx, col_map = find_source_columns(ws_src)
    if not header_row_idx or 'ho_ten' not in col_map:
        stats['errors'].append('Không tìm thấy dòng tiêu đề (phải có cột "Họ và tên") trong file gốc.')
        return None, stats

    # Clear template
    for row_idx in range(3, ws_tpl.max_row + 1):
        for col_idx in range(1, 20):
            ws_tpl.cell(row=row_idx, column=col_idx).value = None

    source_start = header_row_idx + 1
    tpl_start = 3
    row_count = 0
    
    def safe_get(row, key):
        if key in col_map and col_map[key] < len(row):
            return row[col_map[key]]
        return None

    for src_row in ws_src.iter_rows(min_row=source_start, values_only=True):
        ho_ten = safe_get(src_row, 'ho_ten')
        
        # Skip empty names or instruction rows (e.g. "<Họ và tên gồm chữ cái>")
        if not ho_ten or str(ho_ten).strip().startswith('<'):
            continue

        tpl_row = tpl_start + row_count
        stats['total_rows'] += 1

        ws_tpl.cell(row=tpl_row, column=1).value = row_count + 1
        ws_tpl.cell(row=tpl_row, column=2).value = ho_ten

        ngay_sinh_raw = str(safe_get(src_row, 'ngay_sinh') or '').strip()
        if ngay_sinh_raw and len(ngay_sinh_raw) == 4 and ngay_sinh_raw.isdigit():
            ngay_sinh_raw = f'01/01/{ngay_sinh_raw}'
        ws_tpl.cell(row=tpl_row, column=3).value = ngay_sinh_raw

        gt_raw = nfc(safe_get(src_row, 'gioi_tinh'))
        gt_mapped = map_generic(gt_raw, lookups['gioi_tinh'])
        ws_tpl.cell(row=tpl_row, column=4).value = gt_mapped
        if gt_mapped != gt_raw: stats['mapped_gender'] += 1

        qt_raw = nfc(safe_get(src_row, 'quoc_tich'))
        qt_mapped = map_generic(qt_raw, lookups['quoc_tich'])
        ws_tpl.cell(row=tpl_row, column=5).value = qt_mapped
        if qt_mapped != qt_raw: stats['mapped_nationality'] += 1

        # LGT & SGT mapping depending on old vs new format
        if 'loai_giay_to' in col_map and 'so_giay_to' in col_map:
            lgt_raw = nfc(safe_get(src_row, 'loai_giay_to'))
            so_gt_raw = safe_get(src_row, 'so_giay_to')
        else:
            so_cccd = safe_get(src_row, 'so_cccd')
            so_hc = safe_get(src_row, 'so_ho_chieu')
            so_gtk = safe_get(src_row, 'so_giay_to_khac')
            if so_cccd:
                lgt_raw = 'CCCD'
                so_gt_raw = so_cccd
            elif so_hc:
                lgt_raw = 'Hộ chiếu'
                so_gt_raw = so_hc
            elif so_gtk:
                lgt_raw = 'Thẻ BHYT'
                so_gt_raw = so_gtk
            else:
                lgt_raw = ''
                so_gt_raw = ''

        lgt_mapped = map_generic(lgt_raw, lookups['loai_giay_to'])
        ws_tpl.cell(row=tpl_row, column=6).value = lgt_mapped
        if lgt_mapped != lgt_raw: stats['mapped_doc_type'] += 1

        ws_tpl.cell(row=tpl_row, column=7).value = safe_get(src_row, 'ten_giay_to') or ''
        ws_tpl.cell(row=tpl_row, column=8).value = so_gt_raw or ''
        ws_tpl.cell(row=tpl_row, column=9).value = safe_get(src_row, 'so_dt') or ''

        nct_raw = nfc(safe_get(src_row, 'noi_cu_tru'))
        nct_mapped = map_generic(nct_raw, lookups['noi_cu_tru'])
        ws_tpl.cell(row=tpl_row, column=10).value = nct_mapped
        if nct_mapped != nct_raw: stats['mapped_residence'] += 1

        tinh_raw = safe_get(src_row, 'tinh_tp')
        tinh_code = get_tinh_code(tinh_raw, lookups)
        tinh_mapped = map_tinh(tinh_raw, lookups)
        ws_tpl.cell(row=tpl_row, column=11).value = tinh_mapped
        if tinh_mapped != nfc(tinh_raw): stats['mapped_province'] += 1

        px_raw = safe_get(src_row, 'phuong_xa')
        px_mapped = map_phuong_xa(px_raw, tinh_code, lookups)
        ws_tpl.cell(row=tpl_row, column=12).value = px_mapped
        if px_mapped != nfc(px_raw): stats['mapped_ward'] += 1

        ws_tpl.cell(row=tpl_row, column=13).value = safe_get(src_row, 'dia_chi') or ''

        nd_str, nd_dt = fix_date(safe_get(src_row, 'ngay_den'))
        ws_tpl.cell(row=tpl_row, column=14).value = nd_str or ''
        if safe_get(src_row, 'ngay_den') and ' ' in str(safe_get(src_row, 'ngay_den')):
            stats['fixed_date_format'] += 1

        ndi_raw = safe_get(src_row, 'ngay_di')
        if ndi_raw and str(ndi_raw).strip():
            ndi_str, _ = fix_date(ndi_raw)
            ws_tpl.cell(row=tpl_row, column=15).value = ndi_str or ''
        else:
            if nd_dt:
                ngay_di = nd_dt + timedelta(days=extra_days)
                ws_tpl.cell(row=tpl_row, column=15).value = ngay_di.strftime('%d/%m/%Y')
                stats['added_departure_date'] += 1
            else:
                ws_tpl.cell(row=tpl_row, column=15).value = ''

        phong_khoa_raw = str(safe_get(src_row, 'phong_khoa') or '').strip()
        pk_norm = nfc(phong_khoa_raw).upper()
        if "KHOA NGOẠI THẦN KINH" in pk_norm and "CHẤN THƯƠNG CHỈNH HÌNH" in pk_norm:
            phong_khoa_raw = "NGOẠI TK - CTCH"
        ws_tpl.cell(row=tpl_row, column=16).value = phong_khoa_raw

        lydo_raw = nfc(safe_get(src_row, 'ly_do'))
        lydo_mapped = map_generic(lydo_raw, lookups['ly_do'])
        ws_tpl.cell(row=tpl_row, column=17).value = lydo_mapped
        if lydo_mapped != lydo_raw: stats['mapped_reason'] += 1

        row_count += 1

    # Lưu file output vào thư mục outputs (trên server)
    output_id = str(uuid.uuid4())[:8]
    output_filename = f'KBLT_Converted_{output_id}.xlsx'
    output_path = os.path.join(app.config['OUTPUT_FOLDER'], output_filename)
    wb_tpl.save(output_path)

    return output_filename, stats

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/convert', methods=['POST'])
def convert():
    if 'source_file' not in request.files:
        return jsonify({'error': 'Vui lòng tải lên file gốc'}), 400
    source_file = request.files['source_file']
    extra_days = int(request.form.get('extra_days', 3))
    if not source_file.filename:
        return jsonify({'error': 'Vui lòng chọn file gốc'}), 400
    if not os.path.exists(TEMPLATE_PATH):
        return jsonify({'error': 'Không tìm thấy file mẫu KBLT trên server. Vui lòng liên hệ admin.'}), 500

    src_path = os.path.join(app.config['UPLOAD_FOLDER'], f'src_{uuid.uuid4().hex[:8]}_{source_file.filename}')
    source_file.save(src_path)

    try:
        output_filename, stats = process_files(src_path, extra_days)
        if stats['errors']:
            return jsonify({'error': '\n'.join(stats['errors'])}), 400
        return jsonify({
            'success': True,
            'filename': output_filename,
            'stats': stats,
        })
    except Exception as e:
        return jsonify({'error': f'Lỗi xử lý: {str(e)}'}), 500
    finally:
        try: os.remove(src_path)
        except: pass

@app.route('/download/<filename>')
def download_file(filename):
    """Cho phép tải file kết quả về trình duyệt."""
    # Bảo mật: chỉ cho phép tải file KBLT_Converted_*.xlsx
    if not filename.startswith('KBLT_Converted_') or not filename.endswith('.xlsx'):
        return jsonify({'error': 'File không hợp lệ'}), 400
    
    file_path = os.path.join(app.config['OUTPUT_FOLDER'], filename)
    if not os.path.exists(file_path):
        return jsonify({'error': 'File không tồn tại hoặc đã hết hạn. Vui lòng chuyển đổi lại.'}), 404
    
    return send_file(
        file_path,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    print("\n" + "="*60)
    print("  KBLT CONVERTER - THONG MINH PHAT HIEN COT")
    print(f"  File mau: {TEMPLATE_PATH}")
    print(f"  Mo trinh duyet: http://localhost:{port}")
    print("="*60 + "\n")
    app.run(debug=False, host='0.0.0.0', port=port)
